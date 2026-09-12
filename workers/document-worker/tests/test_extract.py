"""모든 fixture는 이 테스트가 즉석에서 생성하는 합성 데이터다 -- 실제 고객
문서나 저장된 바이너리 fixture를 쓰지 않는다 (spec 38절 불변조건)."""

import pymupdf
import openpyxl
from docx import Document

from document_worker.extract import extract, extract_docx, extract_pdf, extract_xlsx


def make_synthetic_pdf(path: str) -> None:
    # ASCII text only: PyMuPDF's built-in base14 fonts have no Korean
    # glyphs, so inserting Korean here would round-trip as mojibake --
    # a font-embedding concern, not something extract_pdf() should be
    # judged on. Real (embedded-font) PDFs extract Korean text fine.
    doc = pymupdf.open()
    page1 = doc.new_page()
    page1.insert_text((72, 72), "first page body text")
    page2 = doc.new_page()
    page2.insert_text((72, 72), "second page body text")
    doc.save(path)
    doc.close()


def make_synthetic_docx(path: str) -> None:
    doc = Document()
    doc.add_paragraph("문단 하나.")
    doc.add_paragraph("문단 둘.")
    doc.save(path)


def make_synthetic_xlsx(path: str) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "매출"
    ws1["A1"] = "연도"
    ws1["B1"] = "매출액"
    ws1["A2"] = 2026
    ws1["B2"] = 1000000
    ws2 = wb.create_sheet("비용")
    ws2["A1"] = "항목"
    ws2["B1"] = "금액"
    wb.save(path)


class TestExtractPdf:
    def test_페이지별로_텍스트를_추출한다(self, tmp_path):
        pdf_path = tmp_path / "sample.pdf"
        make_synthetic_pdf(str(pdf_path))

        pages = extract_pdf(str(pdf_path))

        assert len(pages) == 2
        assert pages[0].page_number == 1
        assert "first page" in pages[0].text
        assert pages[1].page_number == 2
        assert "second page" in pages[1].text


class TestExtractDocx:
    def test_문단_텍스트를_하나의_페이지로_추출한다(self, tmp_path):
        docx_path = tmp_path / "sample.docx"
        make_synthetic_docx(str(docx_path))

        pages = extract_docx(str(docx_path))

        assert len(pages) == 1
        assert pages[0].page_number == 1
        assert "문단 하나." in pages[0].text
        assert "문단 둘." in pages[0].text


class TestExtractXlsx:
    def test_시트마다_하나의_페이지로_추출한다(self, tmp_path):
        xlsx_path = tmp_path / "sample.xlsx"
        make_synthetic_xlsx(str(xlsx_path))

        pages = extract_xlsx(str(xlsx_path))

        assert len(pages) == 2
        assert pages[0].page_number == 1
        assert "매출" in pages[0].text
        assert "2026" in pages[0].text
        assert pages[1].page_number == 2


class TestExtractDispatch:
    def test_media_type으로_올바른_추출기를_선택한다(self, tmp_path):
        pdf_path = tmp_path / "a.pdf"
        make_synthetic_pdf(str(pdf_path))

        pages = extract(str(pdf_path), "application/pdf")

        assert len(pages) == 2

    def test_지원하지_않는_media_type은_ValueError를_던진다(self, tmp_path):
        import pytest

        with pytest.raises(ValueError):
            extract(str(tmp_path / "x.png"), "image/png")
