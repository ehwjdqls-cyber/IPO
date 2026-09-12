"""Text extraction for PDF/DOCX/XLSX, per spec section 15 (업로드 데이터
흐름, step 5) and the document_pages schema (section 17): one row per
"page", where "page" means an actual PDF page, the whole document for
DOCX (Word has no fixed page boundaries in its XML -- this is a known
MVP simplification), or one worksheet for XLSX.
"""

from __future__ import annotations

from dataclasses import dataclass

import openpyxl
import pymupdf
from docx import Document

SUPPORTED_MEDIA_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str


def extract_pdf(path: str) -> list[ExtractedPage]:
    doc = pymupdf.open(path)
    try:
        return [
            ExtractedPage(page_number=index + 1, text=page.get_text())
            for index, page in enumerate(doc)
        ]
    finally:
        doc.close()


def extract_docx(path: str) -> list[ExtractedPage]:
    doc = Document(path)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    return [ExtractedPage(page_number=1, text=text)]


def extract_xlsx(path: str) -> list[ExtractedPage]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    pages: list[ExtractedPage] = []
    for index, sheet_name in enumerate(workbook.sheetnames):
        sheet = workbook[sheet_name]
        lines = [f"[{sheet_name}]"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append("\t".join(cells))
        pages.append(ExtractedPage(page_number=index + 1, text="\n".join(lines)))
    return pages


def extract(path: str, media_type: str) -> list[ExtractedPage]:
    if media_type not in SUPPORTED_MEDIA_TYPES:
        raise ValueError(f"지원하지 않는 media_type입니다: {media_type}")
    if media_type == "application/pdf":
        return extract_pdf(path)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_docx(path)
    return extract_xlsx(path)
